from io import BytesIO
import re
from typing import Callable, Literal, TypedDict
from zipfile import BadZipFile

MAX_INPUT_LENGTH = 30_000
MAX_RESUME_FILE_SIZE = 10 * 1024 * 1024

Skill = tuple[str, tuple[str, ...]]

SKILLS: tuple[Skill, ...] = (
    ("Python", ("python",)),
    ("Java", ("java",)),
    ("JavaScript", ("javascript", "js")),
    ("TypeScript", ("typescript", "ts")),
    ("C++", ("c++", "cpp")),
    ("C#", ("c#", "c sharp")),
    ("Go", ("golang",)),
    ("Rust", ("rust",)),
    ("Ruby", ("ruby",)),
    ("PHP", ("php",)),
    ("SQL", ("sql",)),
    ("HTML", ("html",)),
    ("CSS", ("css",)),
    ("React", ("react", "react.js", "reactjs")),
    ("Angular", ("angular",)),
    ("Vue", ("vue", "vue.js", "vuejs")),
    ("Node.js", ("node.js", "nodejs")),
    ("Express.js", ("express.js", "express")),
    ("Django", ("django",)),
    ("Flask", ("flask",)),
    ("FastAPI", ("fastapi", "fast api")),
    ("Spring", ("spring framework", "spring boot")),
    (".NET", (".net", "dotnet")),
    ("ASP.NET", ("asp.net",)),
    ("REST APIs", ("rest api", "rest apis", "restful api", "restful apis")),
    ("GraphQL", ("graphql",)),
    ("Microservices", ("microservice", "microservices")),
    ("AWS", ("aws", "amazon web services")),
    ("Azure", ("azure", "microsoft azure")),
    ("Google Cloud", ("google cloud", "gcp")),
    ("Docker", ("docker",)),
    ("Kubernetes", ("kubernetes", "k8s")),
    ("Terraform", ("terraform",)),
    ("Linux", ("linux",)),
    ("CI/CD", ("ci/cd", "continuous integration", "continuous delivery")),
    ("Git", ("git",)),
    ("GitHub Actions", ("github actions",)),
    ("Jenkins", ("jenkins",)),
    ("PostgreSQL", ("postgresql", "postgres")),
    ("MySQL", ("mysql",)),
    ("SQL Server", ("sql server", "microsoft sql server")),
    ("MongoDB", ("mongodb", "mongo db")),
    ("Redis", ("redis",)),
    ("Snowflake", ("snowflake",)),
    ("Databricks", ("databricks",)),
    ("Apache Spark", ("apache spark", "pyspark", "spark")),
    ("Apache Kafka", ("apache kafka", "kafka")),
    ("Apache Airflow", ("apache airflow", "airflow")),
    ("Pandas", ("pandas",)),
    ("NumPy", ("numpy",)),
    ("scikit-learn", ("scikit-learn", "scikit learn", "sklearn")),
    ("TensorFlow", ("tensorflow",)),
    ("PyTorch", ("pytorch",)),
    ("Machine learning", ("machine learning",)),
    ("Deep learning", ("deep learning",)),
    ("Natural language processing", ("natural language processing", "nlp")),
    ("Computer vision", ("computer vision",)),
    ("Generative AI", ("generative ai", "genai")),
    ("Large language models", ("large language model", "large language models", "llm", "llms")),
    ("Prompt engineering", ("prompt engineering",)),
    ("Data analysis", ("data analysis", "data analytics")),
    ("Data visualization", ("data visualization", "data viz")),
    ("Power BI", ("power bi",)),
    ("Tableau", ("tableau",)),
    ("Microsoft Excel", ("microsoft excel", "excel")),
    ("Statistics", ("statistics", "statistical analysis")),
    ("A/B testing", ("a/b testing", "ab testing", "split testing")),
    ("Project management", ("project management",)),
    ("Agile", ("agile",)),
    ("Scrum", ("scrum",)),
    ("Jira", ("jira",)),
    ("Communication", ("communication", "written communication", "verbal communication")),
    ("Leadership", ("leadership", "team leadership")),
    ("Customer service", ("customer service", "customer support")),
    ("Sales", ("sales",)),
    ("Marketing", ("marketing",)),
    ("Search engine optimization", ("search engine optimization", "seo")),
    ("Financial modeling", ("financial modeling", "financial modelling")),
    ("Accounting", ("accounting",)),
    ("Payroll", ("payroll",)),
    ("Recruiting", ("recruiting", "recruitment")),
    ("Talent acquisition", ("talent acquisition",)),
    ("UX/UI design", ("ux/ui", "ux design", "ui design", "user experience design")),
    ("Figma", ("figma",)),
    ("User research", ("user research",)),
    ("Product management", ("product management",)),
    ("Software testing", ("software testing", "quality assurance", "qa testing")),
    ("Selenium", ("selenium",)),
    ("Playwright", ("playwright",)),
    ("Cybersecurity", ("cybersecurity", "cyber security")),
    ("Network security", ("network security",)),
    ("Incident response", ("incident response",)),
)


class ResumeAnalysis(TypedDict):
    matching_score: int
    matching_skills: list[str]
    missing_skills: list[str]
    recommendation: Literal["apply", "pass"]
    preparation_tips: list[str]


def extract_resume_text(file_name: str, file_content: bytes) -> str:
    """Extract text from a PDF, DOCX, or plain-text resume upload."""
    if not file_content:
        raise ValueError("The uploaded resume file is empty.")
    if len(file_content) > MAX_RESUME_FILE_SIZE:
        raise ValueError("Resume files must be 10 MB or smaller.")

    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if extension == "txt":
        try:
            text = file_content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("The text file must use UTF-8 encoding.") from error
    elif extension == "pdf":
        try:
            from pypdf import PdfReader
            from pypdf.errors import PdfReadError
        except ImportError as error:
            raise RuntimeError("PDF uploads require the pypdf package.") from error

        try:
            reader = PdfReader(BytesIO(file_content))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDF resumes are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except PdfReadError as error:
            raise ValueError("The uploaded PDF could not be read.") from error
    elif extension == "docx":
        try:
            from docx import Document
            from docx.opc.exceptions import PackageNotFoundError
        except ImportError as error:
            raise RuntimeError("Word uploads require the python-docx package.") from error

        try:
            document = Document(BytesIO(file_content))
        except (PackageNotFoundError, BadZipFile) as error:
            raise ValueError("The uploaded Word document could not be read.") from error

        text_parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            text_parts.extend(cell.text for row in table.rows for cell in row.cells)
        text = "\n".join(text_parts)
    else:
        raise ValueError("Upload a PDF, DOCX, or TXT resume file.")

    if not text.strip():
        if extension == "pdf":
            raise ValueError(
                "No selectable text was found in the PDF. Scanned PDFs need OCR "
                "before they can be analyzed."
            )
        raise ValueError("No text was found in the uploaded resume.")
    return text


def validate_analysis_inputs(resume: str, job_description: str) -> tuple[str, str]:
    """Trim and validate the resume and job description."""
    resume = resume.strip()
    job_description = job_description.strip()
    if not resume:
        raise ValueError("Resume text cannot be empty.")
    if not job_description:
        raise ValueError("Job description cannot be empty.")
    if len(resume) > MAX_INPUT_LENGTH:
        raise ValueError(f"Resume text must be {MAX_INPUT_LENGTH:,} characters or fewer.")
    if len(job_description) > MAX_INPUT_LENGTH:
        raise ValueError(
            f"Job description must be {MAX_INPUT_LENGTH:,} characters or fewer."
        )
    return resume, job_description


def analyze_resume(
    resume: str,
    job_description: str,
    progress_callback: Callable[[str], None] | None = None,
) -> ResumeAnalysis:
    """Compare recognized job skills locally without sending text to a service."""
    resume, job_description = validate_analysis_inputs(resume, job_description)
    if progress_callback:
        progress_callback("Finding recognized skills in the job description...")

    requirements = _find_skills(job_description)
    if not requirements:
        raise ValueError(
            "No supported skills were recognized in the job description. Add "
            "specific skill names or try a description containing technical, "
            "business, or role skills supported by the matcher."
        )

    resume_skills = _find_skills(resume)
    if progress_callback:
        progress_callback("Comparing resume skills with job requirements...")

    matching_skills = [skill for skill in requirements if skill in resume_skills]
    missing_skills = [skill for skill in requirements if skill not in resume_skills]
    weights = {
        skill: _requirement_weight(job_description, skill)
        for skill in requirements
    }
    total_weight = sum(weights.values())
    score = round(sum(weights[skill] for skill in matching_skills) / total_weight * 100)

    return {
        "matching_score": score,
        "matching_skills": matching_skills,
        "missing_skills": missing_skills,
        "recommendation": "apply" if score >= 50 else "pass",
        "preparation_tips": _preparation_tips(missing_skills, matching_skills),
    }


def _find_skills(text: str) -> list[str]:
    found: list[tuple[int, str]] = []
    for label, aliases in SKILLS:
        positions = [
            match.start()
            for alias in aliases
            if (match := _skill_pattern(alias).search(text))
        ]
        if positions:
            found.append((min(positions), label))
    return [label for _, label in sorted(found)]


def _skill_pattern(alias: str) -> re.Pattern[str]:
    return re.compile(rf"(?<!\w){re.escape(alias)}(?!\w)", re.IGNORECASE)


def _requirement_weight(job_description: str, skill: str) -> float:
    required = ("required", "must have", "must be", "need to", "minimum")
    preferred = ("preferred", "nice to have", "bonus", "ideally", "desirable")
    aliases = next((aliases for label, aliases in SKILLS if label == skill), ())
    sentences = re.split(r"(?<=[.!?;])\s+|\n+", job_description)
    context = next(
        (
            sentence.lower()
            for sentence in sentences
            if any(_skill_pattern(alias).search(sentence) for alias in aliases)
        ),
        "",
    )
    if any(marker in context for marker in preferred):
        return 0.5
    if any(marker in context for marker in required):
        return 1.5
    return 1.0


def _preparation_tips(
    missing_skills: list[str],
    matching_skills: list[str],
) -> list[str]:
    tips: list[str] = []
    for skill in missing_skills[:1]:
        tips.append(
            f"Review the role's {skill} requirement. If you have relevant "
            "experience, add a truthful, specific example to your resume."
        )
    if matching_skills:
        strongest = ", ".join(matching_skills[:2])
        tips.append(
            f"Prepare a concise example showing how you used {strongest} and "
            "what result you achieved."
        )
    else:
        tips.append(
            "Review the role requirements and prepare honest examples that show "
            "how your experience transfers to them."
        )
    tips.append(
        "Tailor your resume summary and most relevant experience to this role, "
        "using only claims you can support."
    )
    return tips
